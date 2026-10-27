"""세금계산서 PDF 파일명 변환기

사용법:
  python rename_invoices.py              # 미리보기 (이름만 보여주고 바꾸지 않음)
  python rename_invoices.py --apply      # 실제로 이름 변경 + 월별 폴더 이동
  python rename_invoices.py --config my.json

동작:
  1. inbox 폴더의 PDF에서 승인번호·작성일자·거래처·공급가액을 읽는다.
  2. 홈택스 목록 엑셀이 있으면 승인번호로 대조해 그 값을 우선 쓴다.
  3. {date}_{partner}_{amount}_{kind}.pdf 로 이름을 바꿔 outbox/YYYY-MM/매입|매출/ 로 옮긴다.
  4. 못 읽은 파일은 outbox/확인필요/ 로 옮기고, 읽은 글자를 .txt로 같이 둔다.
  5. 처리 결과를 outbox/처리결과_YYYYMMDD_HHMM.xlsx 에 남긴다.
"""
import argparse, json, re, shutil, sys
from datetime import datetime
from pathlib import Path

import pdfplumber
from openpyxl import Workbook, load_workbook

APPROVAL_RE = re.compile(r"\d{8}-\d{8}-\d{8}")
DATE_RE = re.compile(r"(20\d{2})[.\-/년 ]\s*(\d{1,2})[.\-/월 ]\s*(\d{1,2})일?")
BIZNO_RE = re.compile(r"\d{3}-\d{2}-\d{5}")
AMOUNT_RE = re.compile(r"[\d,]{4,}")
ILLEGAL = re.compile(r'[\\/:*?"<>|\r\n\t]')


def load_config(path):
    cfg = json.loads(Path(path).read_text(encoding="utf-8"))
    base = Path(path).resolve().parent
    for k in ("inbox", "outbox", "hometax_list"):
        if cfg.get(k):
            cfg[k] = (base / cfg[k]).resolve()
    cfg["our_business_no"] = cfg.get("our_business_no", "").replace(" ", "")
    return cfg


def clean_partner(name):
    name = re.sub(r"\(주\)|㈜|주식회사|유한회사|\(유\)", "", name or "")
    name = ILLEGAL.sub("", name).strip(" _-")
    return name[:30] or "거래처미상"


def to_int(s):
    s = re.sub(r"[^\d]", "", str(s or ""))
    return int(s) if s else None


def norm_date(s):
    if isinstance(s, datetime):
        return s.strftime("%Y%m%d")
    m = DATE_RE.search(str(s or ""))
    return f"{m.group(1)}{int(m.group(2)):02d}{int(m.group(3)):02d}" if m else None


# ---------- 홈택스 목록 엑셀 ----------
def load_hometax_list(path):
    """승인번호 -> {date, partner, amount, kind} 사전. 열 이름은 '포함' 기준으로 찾는다."""
    if not path or not Path(path).exists():
        return {}
    ws = load_workbook(path, read_only=True, data_only=True).active
    rows = ws.iter_rows(values_only=True)
    header = None
    for r in rows:
        if r and any(c and "승인번호" in str(c) for c in r):
            header = [str(c or "").replace(" ", "") for c in r]
            break
    if not header:
        return {}

    def col(*keys):
        for i, h in enumerate(header):
            if all(k in h for k in keys):
                return i
        return None

    i_no, i_date, i_amt = col("승인번호"), col("작성일자"), col("공급가액")
    i_sup, i_buy, i_kind = col("공급자", "상호"), col("공급받는자", "상호"), col("구분")
    out = {}
    for r in rows:
        if not r or i_no is None or not r[i_no]:
            continue
        no = str(r[i_no]).strip()
        kind = None
        if i_kind is not None and r[i_kind]:
            kind = "매출" if "매출" in str(r[i_kind]) else "매입" if "매입" in str(r[i_kind]) else None
        sup = r[i_sup] if i_sup is not None else None
        buy = r[i_buy] if i_buy is not None else None
        partner = buy if kind == "매출" else sup if kind == "매입" else (sup or buy)
        out[no] = {
            "date": norm_date(r[i_date]) if i_date is not None else None,
            "partner": clean_partner(partner),
            "amount": to_int(r[i_amt]) if i_amt is not None else None,
            "kind": kind,
            "source": "홈택스목록",
        }
    return out


# ---------- PDF 본문 ----------
def parse_pdf(path, our_no):
    with pdfplumber.open(path) as pdf:
        text = "\n".join((p.extract_text() or "") for p in pdf.pages)
    info = {"text": text, "approval": None, "date": None, "partner": None, "amount": None, "kind": None, "source": "PDF"}
    m = APPROVAL_RE.search(text)
    info["approval"] = m.group(0) if m else None

    m = re.search(r"작성일자[^\d]{0,10}(20\d{2}[.\-/년 ]\s*\d{1,2}[.\-/월 ]\s*\d{1,2})", text)
    info["date"] = norm_date(m.group(1)) if m else norm_date(text)

    # 공급자 / 공급받는자 블록 분리: '공급받는자'가 먼저 나오는 줄 이후를 받는자 블록으로 본다
    sup_block, buy_block = text, ""
    m = re.search(r"공\s*급\s*받\s*는\s*자", text)
    if m:
        sup_block, buy_block = text[: m.start()], text[m.start():]

    def biz(block):
        m = BIZNO_RE.search(block)
        return m.group(0) if m else None

    def company(block):
        m = re.search(r"상\s*호(?:\s*\(법인명\))?\s*[:：]?\s*([^\n]+?)(?:\s+성\s*명|\s+대표자|\n|$)", block)
        return m.group(1).strip() if m else None

    sup_no, buy_no = biz(sup_block), biz(buy_block)
    if our_no and sup_no == our_no:
        info["kind"], info["partner"] = "매출", company(buy_block)
    elif our_no and buy_no == our_no:
        info["kind"], info["partner"] = "매입", company(sup_block)
    else:
        info["partner"] = company(sup_block) or company(buy_block)

    m = re.search(r"공\s*급\s*가\s*액[^\d]{0,20}([\d,]{4,})", text)
    info["amount"] = to_int(m.group(1)) if m else None
    info["partner"] = clean_partner(info["partner"]) if info["partner"] else None
    return info


def build_name(cfg, rec):
    amount = f"{rec['amount']:,}" if rec.get("amount") is not None else "금액미상"
    return cfg["name_pattern"].format(
        date=rec.get("date") or "날짜미상",
        partner=rec.get("partner") or "거래처미상",
        amount=amount,
        kind=rec.get("kind") or "구분미상",
    )


def unique(path):
    if not path.exists():
        return path
    for i in range(2, 100):
        p = path.with_name(f"{path.stem}_{i}{path.suffix}")
        if not p.exists():
            return p
    raise RuntimeError(f"같은 이름이 너무 많음: {path}")


def process(cfg, apply):
    inbox, outbox = Path(cfg["inbox"]), Path(cfg["outbox"])
    listing = load_hometax_list(cfg.get("hometax_list"))
    results = []
    for pdf in sorted(inbox.glob("*.pdf")):
        try:
            rec = parse_pdf(pdf, cfg["our_business_no"])
        except Exception as e:  # 깨진 PDF 등
            rec = {"text": f"읽기 실패: {e}", "approval": None, "source": "PDF"}
        if rec.get("approval") in listing:
            rec = {**rec, **listing[rec["approval"]]}
        ok = all(rec.get(k) for k in ("date", "partner", "kind")) and rec.get("amount") is not None
        new_name = build_name(cfg, rec) + ".pdf"
        if ok:
            sub = outbox / (f"{rec['date'][:4]}-{rec['date'][4:6]}" if cfg.get("month_folder") else "") / rec["kind"]
        else:
            sub = outbox / "확인필요"
        dest = unique(sub / new_name)
        if apply:
            sub.mkdir(parents=True, exist_ok=True)
            shutil.move(str(pdf), dest)
            if not ok:
                dest.with_suffix(".txt").write_text(rec.get("text", ""), encoding="utf-8")
        results.append({
            "원본": pdf.name, "새이름": str(dest.relative_to(outbox)), "상태": "완료" if ok else "확인필요",
            "작성일자": rec.get("date"), "거래처": rec.get("partner"), "공급가액": rec.get("amount"),
            "구분": rec.get("kind"), "승인번호": rec.get("approval"), "출처": rec.get("source"),
        })
    return results


def write_log(outbox, results):
    outbox.mkdir(parents=True, exist_ok=True)
    wb = Workbook(); ws = wb.active; ws.title = "처리결과"
    cols = ["원본", "새이름", "상태", "작성일자", "거래처", "공급가액", "구분", "승인번호", "출처"]
    ws.append(cols)
    for r in results:
        ws.append([r.get(c) for c in cols])
    p = outbox / f"처리결과_{datetime.now():%Y%m%d_%H%M}.xlsx"
    wb.save(p)
    return p


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--config", default=str(Path(__file__).with_name("config.json")))
    ap.add_argument("--apply", action="store_true", help="실제로 이름을 바꾸고 이동 (없으면 미리보기)")
    a = ap.parse_args()
    cfg = load_config(a.config)
    results = process(cfg, a.apply)
    if not results:
        print(f"PDF 없음: {cfg['inbox']}"); return
    w = max(len(r["원본"]) for r in results)
    for r in results:
        print(f"[{r['상태']}] {r['원본']:<{w}} -> {r['새이름']}")
    done = sum(r["상태"] == "완료" for r in results)
    print(f"\n{'적용' if a.apply else '미리보기'}: 총 {len(results)}건, 완료 {done}건, 확인필요 {len(results)-done}건")
    if a.apply:
        print(f"결과 파일: {write_log(Path(cfg['outbox']), results)}")
    else:
        print("실제로 바꾸려면 --apply 를 붙여 실행하세요.")


if __name__ == "__main__":
    main()
