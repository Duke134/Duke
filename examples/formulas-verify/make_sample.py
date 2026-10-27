"""샘플 런웨이·급여 워크북 생성 (가상 숫자).

- sample_clean.xlsx : 정상 시트
- sample_broken.xlsx: 실무에서 자주 나는 실수 3개를 심은 시트
    1) 런웨이!H7  수식 자리에 값 하드코딩 (붙여넣기 실수)
    2) 런웨이!C9  성장률 수식이 전월(C8) 대신 C7을 참조
    3) 급여!B6    합계 SUM 범위가 마지막 직원 행을 빠뜨림
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")  # 입력 셀은 노란색
BOLD = Font(bold=True)


def build(path: str, broken: bool = False) -> None:
    wb = Workbook()

    # ── 입력 시트 ──────────────────────────────────────────────
    ws = wb.active
    ws.title = "입력"
    rows = [
        ("항목", "값", "메모"),
        ("기초현금(2026-10말)", 1_200_000_000, "가상 숫자"),
        ("월 매출", 60_000_000, "가상 숫자"),
        ("월 인건비", 85_000_000, "가상 숫자"),
        ("월 임차료", 12_000_000, "가상 숫자"),
        ("월 기타비용", 18_000_000, "가상 숫자"),
        ("월 매출 성장률", 0.03, "가정"),
    ]
    for r, row in enumerate(rows, start=1):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            if r == 1:
                cell.font = BOLD
            elif c == 2:
                cell.fill = INPUT_FILL
    ws["B7"].number_format = "0.0%"
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 18

    # ── 런웨이 시트 (12개월) ───────────────────────────────────
    rw = wb.create_sheet("런웨이")
    header = ["월", "기초현금", "매출", "인건비", "임차료", "기타비용", "순현금흐름", "기말현금"]
    for c, h in enumerate(header, start=1):
        rw.cell(row=1, column=c, value=h).font = BOLD
    months = [f"2026-{m:02d}" for m in (11, 12)] + [f"2027-{m:02d}" for m in range(1, 11)]
    for i, m in enumerate(months):
        r = i + 2
        rw[f"A{r}"] = m
        rw[f"B{r}"] = "=입력!$B$2" if r == 2 else f"=H{r-1}"
        rw[f"C{r}"] = "=입력!$B$3" if r == 2 else f"=C{r-1}*(1+입력!$B$7)"
        rw[f"D{r}"] = "=입력!$B$4"
        rw[f"E{r}"] = "=입력!$B$5"
        rw[f"F{r}"] = "=입력!$B$6"
        rw[f"G{r}"] = f"=C{r}-D{r}-E{r}-F{r}"
        rw[f"H{r}"] = f"=B{r}+G{r}"
    rw["J1"] = "요약"
    rw["J1"].font = BOLD
    rw["J2"] = "평균 월 순현금흐름"
    rw["K2"] = "=AVERAGE(G2:G13)"
    rw["J3"] = "런웨이(개월, 기초현금/평균소진)"
    rw["K3"] = "=IF(K2<0,입력!B2/ABS(K2),\"흑자\")"
    rw["J4"] = "12개월 내 현금 플러스 유지 월수"
    rw["K4"] = "=COUNTIF(H2:H13,\">0\")"
    for col in "BCDEFGH":
        for r in range(2, 14):
            rw[f"{col}{r}"].number_format = "#,##0"
    rw["K2"].number_format = "#,##0"
    rw["K3"].number_format = "0.0"
    rw.column_dimensions["J"].width = 32
    for col in "BCDEFGHK":
        rw.column_dimensions[col].width = 15

    if broken:
        rw["H7"] = 996_000_000            # 실수 1: 하드코딩 값
        rw["C9"] = "=C7*(1+입력!$B$7)"     # 실수 2: 잘못된 행 참조

    # ── 급여 시트 (가상 3명) ──────────────────────────────────
    pw = wb.create_sheet("급여")
    pw["F1"] = "공제 요율(가정값, 실제 적용 전 확인 필요)"
    pw["F1"].font = BOLD
    rates = [("국민연금(근로자)", 0.045), ("건강보험(근로자)", 0.03545),
             ("장기요양(건강보험료 대비)", 0.1295), ("고용보험(근로자)", 0.009)]
    for i, (k, v) in enumerate(rates, start=2):
        pw[f"F{i}"] = k
        pw[f"G{i}"] = v
        pw[f"G{i}"].fill = INPUT_FILL
        pw[f"G{i}"].number_format = "0.000%"
    pw.column_dimensions["F"].width = 28

    hdr = ["이름", "기본급", "식대(비과세)", "지급총액", "국민연금", "건강보험", "장기요양", "고용보험", "공제합계", "실지급액"]
    # A~D 그리고 H~M 에 배치 (E 비움, F/G는 요율)
    cols = ["A", "B", "C", "D", "H", "I", "J", "K", "L", "M"]
    for col, h in zip(cols, hdr):
        pw[f"{col}1"] = h
        pw[f"{col}1"].font = BOLD
    staff = [("직원A", 4_200_000), ("직원B", 3_600_000), ("직원C", 3_000_000)]
    for i, (name, base) in enumerate(staff, start=2):
        pw[f"A{i}"] = name
        pw[f"B{i}"] = base
        pw[f"B{i}"].fill = INPUT_FILL
        pw[f"C{i}"] = 200_000
        pw[f"C{i}"].fill = INPUT_FILL
        pw[f"D{i}"] = f"=B{i}+C{i}"
        pw[f"H{i}"] = f"=ROUNDDOWN(B{i}*$G$2,-1)"
        pw[f"I{i}"] = f"=ROUNDDOWN(B{i}*$G$3,-1)"
        pw[f"J{i}"] = f"=ROUNDDOWN(I{i}*$G$4,-1)"
        pw[f"K{i}"] = f"=ROUNDDOWN(B{i}*$G$5,-1)"
        pw[f"L{i}"] = f"=SUM(H{i}:K{i})"
        pw[f"M{i}"] = f"=D{i}-L{i}"
    pw["A6"] = "합계"
    pw["A6"].font = BOLD
    for col in ["B", "C", "D", "H", "I", "J", "K", "L", "M"]:
        pw[f"{col}6"] = f"=SUM({col}2:{col}4)"
    if broken:
        pw["B6"] = "=SUM(B2:B3)"            # 실수 3: 합계 범위 누락
    for col in ["B", "C", "D", "H", "I", "J", "K", "L", "M"]:
        for r in range(2, 7):
            pw[f"{col}{r}"].number_format = "#,##0"
        pw.column_dimensions[col].width = 13

    wb.save(path)
    print("saved", path)


if __name__ == "__main__":
    build("sample_clean.xlsx", broken=False)
    build("sample_broken.xlsx", broken=True)
