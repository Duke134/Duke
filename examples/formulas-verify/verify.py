"""엑셀 없이 런웨이·급여 시트를 재검산하는 예시.

사용법
  python verify.py sample_clean.xlsx
  python verify.py sample_broken.xlsx
  python verify.py sample_clean.xlsx --set "입력!B4=93500000"     # 인건비 10% 증가 시나리오

하는 일
  1. formulas 라이브러리로 워크북 전체를 재계산 (Excel 불필요)
  2. 구조 점검: 수식 열의 하드코딩 값 / 행 간 수식 불일치 / 합계 SUM 범위 누락
  3. 독립 재계산: 입력 셀만 읽어 파이썬으로 다시 계산 → formulas 결과와 대조
  4. 4필드(결과·계산과정·근거·주의사항) 보고 + ✅ 검증완료 / ❌ 불일치
"""
import argparse
import re
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from formulas import ExcelModel

TOL = 0.5  # 원 단위 반올림 허용


# ──────────────────────────── 유틸 ────────────────────────────
def key_of(sol, sheet, addr):
    """formulas 결과 딕셔너리에서 '[파일]시트'!주소 키를 대소문자 무시로 찾는다."""
    target = f"]{sheet}'!{addr}".upper()
    for k in sol:
        if k.upper().endswith(target):
            return k
    return None


def val(sol, sheet, addr):
    k = key_of(sol, sheet, addr)
    if k is None:
        return None
    v = sol[k].value
    try:
        v = v[0][0]
    except Exception:
        pass
    return float(v) if isinstance(v, (int, float)) or hasattr(v, "dtype") else v


def rounddown(x, digits):
    """Excel ROUNDDOWN 재현. Excel은 15자리 유효숫자로 정리한 뒤 절사하므로
    32399.999999999996 같은 부동소수점 잔차가 32400으로 처리된다. 순수 int() 절사는 틀린다."""
    x = float(f"{x:.15g}")
    q = 10 ** (-digits)
    return int(x / q) * q if x >= 0 else -int(-x / q) * q


# ──────────────────────────── 구조 점검 ────────────────────────────
def check_column_consistency(ws, col, first_row, last_row):
    """첫 수식 행을 기준으로 아래 행들의 수식이 같은 패턴인지 확인."""
    issues = []
    base = ws[f"{col}{first_row}"].value
    if not (isinstance(base, str) and base.startswith("=")):
        return [f"{ws.title}!{col}{first_row}: 기준 셀이 수식이 아님 ({base!r})"]
    for r in range(first_row + 1, last_row + 1):
        cell = ws[f"{col}{r}"]
        expected = Translator(base, origin=f"{col}{first_row}").translate_formula(f"{col}{r}")
        actual = cell.value
        if not (isinstance(actual, str) and actual.startswith("=")):
            issues.append(f"{ws.title}!{col}{r}: 수식 열에 값 하드코딩 ({actual!r}) — 기대 수식 {expected}")
        elif actual.replace(" ", "") != expected.replace(" ", ""):
            issues.append(f"{ws.title}!{col}{r}: 수식 불일치 — 실제 {actual} / 기대 {expected}")
    return issues


def check_sum_range(ws, addr, want_first, want_last):
    """합계 셀의 SUM 범위가 데이터 행을 모두 덮는지 확인."""
    f = ws[addr].value
    m = re.fullmatch(r"=SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)", str(f).replace(" ", ""))
    if not m:
        return [f"{ws.title}!{addr}: SUM 수식이 아님 ({f!r})"]
    c1, r1, c2, r2 = m.group(1), int(m.group(2)), m.group(3), int(m.group(4))
    if c1 != c2 or r1 != want_first or r2 != want_last:
        return [f"{ws.title}!{addr}: SUM 범위 {c1}{r1}:{c2}{r2} — 기대 {c1}{want_first}:{c1}{want_last}"]
    return []


def structural_checks(wb):
    issues = []
    rw = wb["런웨이"]
    # 2행은 입력 참조, 3행부터 패턴 반복
    for col in "BC":
        issues += check_column_consistency(rw, col, 3, 13)
    for col in "DEFGH":
        issues += check_column_consistency(rw, col, 2, 13)
    pw = wb["급여"]
    for col in ["D", "H", "I", "J", "K", "L", "M"]:
        issues += check_column_consistency(pw, col, 2, 4)
    for col in ["B", "C", "D", "H", "I", "J", "K", "L", "M"]:
        issues += check_sum_range(pw, f"{col}6", 2, 4)
    return issues


# ──────────────────────────── 독립 재계산 ────────────────────────────
def independent_runway(inputs, months=12):
    cash, sales = inputs["기초현금"], inputs["월매출"]
    flows, ends = [], []
    for i in range(months):
        if i > 0:
            sales *= 1 + inputs["성장률"]
        net = sales - inputs["인건비"] - inputs["임차료"] - inputs["기타"]
        cash += net
        flows.append(net)
        ends.append(cash)
    avg = sum(flows) / months
    runway = inputs["기초현금"] / abs(avg) if avg < 0 else None
    return {"기말현금_12개월": ends[-1], "평균순현금흐름": avg, "런웨이": runway,
            "플러스월수": sum(1 for e in ends if e > 0)}


def independent_payroll(rows, rates):
    out, total_net = [], 0
    for base, meal in rows:
        gross = base + meal
        np_ = rounddown(base * rates["국민연금"], -1)
        hi = rounddown(base * rates["건강"], -1)
        ltc = rounddown(hi * rates["장기요양"], -1)
        ei = rounddown(base * rates["고용"], -1)
        net = gross - (np_ + hi + ltc + ei)
        out.append(net)
        total_net += net
    return out, total_net


# ──────────────────────────── 메인 ────────────────────────────
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--set", action="append", default=[], metavar="시트!셀=값",
                    help="입력 셀 덮어쓰기(시나리오). 예: 입력!B4=93500000")
    args = ap.parse_args()
    path = Path(args.xlsx)
    fname = path.name

    # 1) 수식 원문 + 입력값 읽기 (openpyxl)
    wb_f = load_workbook(path)                 # 수식
    wb_v = load_workbook(path, data_only=True)  # 캐시값(Excel이 저장했을 때만 존재)

    overrides = {}
    for s in args.set:
        ref, v = s.split("=")
        sheet, addr = ref.split("!")
        overrides[(sheet, addr.upper())] = float(v)

    def inp(sheet, addr):
        if (sheet, addr) in overrides:
            return overrides[(sheet, addr)]
        return wb_f[sheet][addr].value

    # 2) formulas 재계산
    xl = ExcelModel().loads(str(path)).finish()
    f_inputs = {f"'[{fname}]{s}'!{a}": v for (s, a), v in overrides.items()}
    sol = xl.calculate(inputs=f_inputs) if f_inputs else xl.calculate()

    # 3) 구조 점검
    struct_issues = structural_checks(wb_f)

    # 4) 독립 재계산과 대조
    mismatches = []
    rin = {"기초현금": inp("입력", "B2"), "월매출": inp("입력", "B3"), "인건비": inp("입력", "B4"),
           "임차료": inp("입력", "B5"), "기타": inp("입력", "B6"), "성장률": inp("입력", "B7")}
    ind = independent_runway(rin)
    f_end = val(sol, "런웨이", "H13")
    f_avg = val(sol, "런웨이", "K2")
    f_run = val(sol, "런웨이", "K3")
    f_pos = val(sol, "런웨이", "K4")
    for name, a, b in [("런웨이!H13 기말현금", f_end, ind["기말현금_12개월"]),
                       ("런웨이!K2 평균순현금흐름", f_avg, ind["평균순현금흐름"]),
                       ("런웨이!K4 플러스월수", f_pos, ind["플러스월수"])]:
        if abs(a - b) > TOL:
            mismatches.append(f"{name}: 시트 {a:,.0f} vs 독립계산 {b:,.0f}")
    if ind["런웨이"] is not None and (not isinstance(f_run, float) or abs(f_run - ind["런웨이"]) > 0.01):
        mismatches.append(f"런웨이!K3 런웨이: 시트 {f_run:.2f} vs 독립계산 {ind['런웨이']:.2f}")

    rates = {"국민연금": inp("급여", "G2"), "건강": inp("급여", "G3"),
             "장기요양": inp("급여", "G4"), "고용": inp("급여", "G5")}
    rows = [(inp("급여", f"B{r}"), inp("급여", f"C{r}")) for r in range(2, 5)]
    nets, total_net = independent_payroll(rows, rates)
    for r, n in zip(range(2, 5), nets):
        a = val(sol, "급여", f"M{r}")
        if abs(a - n) > TOL:
            mismatches.append(f"급여!M{r} 실지급액: 시트 {a:,.0f} vs 독립계산 {n:,.0f}")
    a = val(sol, "급여", "M6")
    if abs(a - total_net) > TOL:
        mismatches.append(f"급여!M6 실지급 합계: 시트 {a:,.0f} vs 독립계산 {total_net:,.0f}")
    a = val(sol, "급여", "B6")
    b = sum(r[0] for r in rows)
    if abs(a - b) > TOL:
        mismatches.append(f"급여!B6 기본급 합계: 시트 {a:,.0f} vs 독립계산 {b:,.0f}")

    # 5) 캐시값 대조 (Excel이 저장한 파일일 때만 의미 있음)
    cache_note = "없음"
    cached = wb_v["런웨이"]["H13"].value
    if cached is None:
        cache_note = "파일에 Excel 캐시값 없음(openpyxl 생성 파일). 재계산 결과만 사용"
    elif not overrides and abs(cached - f_end) > TOL:
        mismatches.append(f"런웨이!H13 Excel 저장값 {cached:,.0f} vs 재계산 {f_end:,.0f} (저장 후 입력 변경 또는 수동 계산 모드 의심)")

    # ──────────────── 보고 ────────────────
    ok = not struct_issues and not mismatches
    runway_txt = f"{f_run:.1f}개월" if isinstance(f_run, float) else str(f_run)
    print(f"\n📄 {fname}" + (f"   시나리오: {', '.join(args.set)}" if args.set else ""))
    print("=" * 64)
    print("결과")
    print(f"  12개월 후 기말현금  {f_end:>18,.0f} 원")
    print(f"  평균 월 순현금흐름  {f_avg:>18,.0f} 원")
    print(f"  런웨이              {runway_txt:>18}")
    print(f"  실지급액 합계(급여) {val(sol, '급여', 'M6'):>18,.0f} 원")
    print("계산과정")
    print(f"  기말현금 = 기초현금 {rin['기초현금']:,.0f} + Σ(매출−인건비−임차료−기타), 매출은 월 {rin['성장률']:.1%} 성장")
    print(f"  런웨이 = 기초현금 / |평균 월 순현금흐름| = {rin['기초현금']:,.0f} / {abs(f_avg):,.0f}")
    print(f"  실지급액 = (기본급+식대) − ROUNDDOWN(기본급×요율, -1) 합계, 요율은 급여!G2:G5")
    print("근거")
    print(f"  {fname} 입력!B2:B7, 급여!G2:G5, 수식 원문 전체 (formulas {__import__('formulas').__version__} 재계산)")
    print("주의사항")
    print(f"  숫자는 전부 가상 샘플. 공제 요율은 시트 가정값이며 실제 적용 전 확인 필요. 캐시값: {cache_note}")
    print("-" * 64)
    if struct_issues:
        print("구조 점검 ❌")
        for s in struct_issues:
            print("  -", s)
    else:
        print("구조 점검 ✅ 하드코딩·수식 불일치·SUM 범위 누락 없음")
    if mismatches:
        print("독립 재계산 대조 ❌")
        for m in mismatches:
            print("  -", m)
    else:
        print("독립 재계산 대조 ✅ 모든 대조 항목 일치 (허용오차 ±0.5원)")
    print("=" * 64)
    print("✅ 검증완료" if ok else "❌ 검증 실패 — 위 항목 수정 전에는 '검증완료'를 붙이지 않음")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
